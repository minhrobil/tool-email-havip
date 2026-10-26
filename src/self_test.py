"""Kiểm tra thư viện trong bản đóng gói bằng dữ liệu tạm, không cần đăng nhập."""
from pathlib import Path
from tempfile import TemporaryDirectory


def verify_pdf_excel() -> None:
    # Import tại đây để phát hiện cả lỗi thiếu tài nguyên khi nạp thư viện.
    import fitz
    import openpyxl

    from .excel.writer import ExcelWriter
    from .parser.rules import parse_document

    with TemporaryDirectory() as work:
        folder = Path(work)
        pdf = folder / "sample.pdf"
        with fitz.open() as document:
            page = document.new_page()
            page.insert_text((72, 72), (
                "So: 12345/SHTT-NH.IP\n"
                "This sample checks PDF text extraction in the packaged application."
            ))
            document.save(pdf)

        parsed = parse_document(pdf_path=pdf)
        if parsed.so_cong_van_num != "12345":
            raise RuntimeError("Không đọc được số công văn từ PDF mẫu.")

        writer = ExcelWriter(folder, "sample.xlsx")
        writer.append_data_row({
            "Ngày nhận công văn": 1,
            "Số công văn": parsed.so_cong_van_num,
        })
        writer.append_meta_row({"message_id": "self-test"})
        workbook = openpyxl.load_workbook(writer.excel_path)
        try:
            if workbook["DATA"].cell(2, 2).value != "12345":
                raise RuntimeError("Dữ liệu Excel mẫu không đúng.")
            if workbook["META"].cell(2, 1).value != "self-test":
                raise RuntimeError("Thiếu dữ liệu META trong Excel mẫu.")
        finally:
            workbook.close()
